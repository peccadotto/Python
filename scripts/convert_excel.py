### # CONVERT EXCEL TO CSV FUNCTION
### # convert_excel(file_path, output_dir = "output_csv", mode = "sheets", delimiter = ";")

import os
import pandas as pd
import openpyxl

def clean_filename(name):
    return "".join(c for c in name if c.isalnum() or c in (' ', '_', '-')).strip()

def convert_excel(file_path, output_dir="output_csv", mode = "sheets", delimiter = ";"):
    """
    Converts an Excel file into CSV format
    
    Parameters:
    - file_path: Path to the Excel file (.xlsx, .xls)
        e.g. "test.xls", "test.xlsx", "test/test.xls"
    - output_dir: destination directory for the CSV files
        default value = "output_csv"
    - mode: 
        'sheets' -> Converts EVERY sheet into a separate CSV file
        'single' -> Converts only the FIRST sheet into a single CSV
        'tables' -> Converts EVERY official Excel Table (ListObject) into a separate CSV file
    - delimiter: CSV delimiter
        default value = ";"
    """
    if not os.path.exists(file_path):
        print(f"Error: The file '{file_path}' does not exist")
        return

    os.makedirs(output_dir, exist_ok=True)
    base_name = os.path.splitext(os.path.basename(file_path))[0]

    # -------------------------------------------------------------------------
    # CASE 1: Convert the first sheet of the file
    # -------------------------------------------------------------------------
    if mode == "single":
        print("Selected mode: single sheet -> CSV file")
        df = pd.read_excel(file_path)
        out_file = os.path.join(output_dir, f"{base_name}.csv")
        df.to_csv(out_file, index=False, sep=delimiter, encoding='utf-8-sig')
        print(f"[✓] File '{out_file}' saved")

    # -------------------------------------------------------------------------
    # CASE 2: Convert EVERY sheet into a separate CSV file
    # -------------------------------------------------------------------------
    elif mode == "sheets":
        print("Selected mode: every sheet -> separate CSV files")
        xls = pd.ExcelFile(file_path)
        
        for sheet_name in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet_name)
            safe_sheet = clean_filename(sheet_name)
            out_file = os.path.join(output_dir, f"{base_name}_{safe_sheet}.csv")
            df.to_csv(out_file, index=False, sep=delimiter, encoding='utf-8-sig')
            print(f"[✓] Sheet '{sheet_name}' saved")

    # -------------------------------------------------------------------------
    # CASE 3: Convert EVERY Excel Table into a separate CSV file
    # -------------------------------------------------------------------------
    elif mode == "tables":
        print("Selected mode: every table -> separate CSV files")
        wb = openpyxl.load_workbook(file_path, data_only=True)
        table_count = 0

        for ws in wb.worksheets:
            for table_name in ws.tables:
                table_obj = ws.tables[table_name]
                cell_range = table_obj.ref if hasattr(table_obj, 'ref') else ws.tables[table_name]
                
                data = ws[cell_range]
                
                # Extract headers and data rows
                headers = [cell.value for cell in data[0]]
                rows = [[cell.value for cell in row] for row in data[1:]]
                
                df = pd.DataFrame(rows, columns=headers)
                safe_sheet = clean_filename(ws.title)
                safe_table = clean_filename(table_name)
                
                out_file = os.path.join(output_dir, f"{base_name}_{safe_sheet}_{safe_table}.csv")
                df.to_csv(out_file, index=False, sep=delimiter, encoding='utf-8-sig')
                table_count += 1
                print(f"[✓] Table '{table_name}' saved")

        if table_count == 0:
            print("No formatted Excel tables (ListObject) found. Try using mode = 'sheets'")